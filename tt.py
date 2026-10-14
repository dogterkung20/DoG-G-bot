import os
import re
import json
import threading
from datetime import datetime, timedelta
import openpyxl
from flask import Flask
from PIL import Image, ImageChops
from telegram import Update, ReplyKeyboardMarkup, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    ApplicationBuilder,
    ContextTypes,
    CommandHandler,
    MessageHandler,
    ConversationHandler,
    CallbackQueryHandler,
    filters,
)
from telegram.request import HTTPXRequest

# ==========================================
# 🔑 Token บอท Telegram
# ==========================================
TELEGRAM_BOT_TOKEN = "8863260350:AAECdIo15nDl0kxPR6tCnBys9opMExapTY4"
DB_FILE = "local_database.json"

def init_db():
    if not os.path.exists(DB_FILE):
        data = {"customers": [], "documents": [], "last_doc_number": "WN-P56-52"}
        with open(DB_FILE, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=4)

def load_db():
    init_db()
    with open(DB_FILE, 'r', encoding='utf-8') as f:
        return json.load(f)

def save_db(data):
    with open(DB_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=4)

def save_or_update_customer(cust_data):
    db = load_db()
    customers = db.get("customers", [])
    existing_index = -1
    for idx, c in enumerate(customers):
        if (c.get("customer_name") and c.get("customer_name").strip() == cust_data.get("customer_name", "").strip()) or \
           (c.get("tax_id") and c.get("tax_id") != "-" and c.get("tax_id") == cust_data.get("tax_id")):
            existing_index = idx
            break

    if existing_index >= 0:
        customers[existing_index].update(cust_data)
    else:
        customers.append(cust_data)

    db["customers"] = customers
    save_db(db)

def get_next_doc_number():
    db = load_db()
    last_num = db.get("last_doc_number", "WN-P56-52")
    match = re.search(r'(\d+)$', last_num)
    if match:
        next_val = int(match.group(1)) + 1
        prefix = last_num[:match.start()]
        new_num = f"{prefix}{next_val:02d}"
    else:
        new_num = "WN-P56-53"
    db["last_doc_number"] = new_num
    save_db(db)
    return new_num

def get_thai_dates():
    now = datetime.now()
    thai_year = now.year + 543
    date_str = f"{now.day:02d}/{now.month:02d}/{thai_year}"
    due_date = now + timedelta(days=7)
    due_thai_year = due_date.year + 543
    due_date_str = f"{due_date.day:02d}/{due_date.month:02d}/{due_thai_year}"
    return date_str, due_date_str, now.strftime("%Y-%m-%d")

# --- PyMuPDF: ใช้ตัดขอบขาวของ PDF และแปลงเป็น PNG ---
try:
    import fitz  # PyMuPDF
except ImportError:
    try:
        import pymupdf as fitz
    except ImportError:
        fitz = None

A4_WIDTH_PT = 595.276
CONVERT_ERRORS = {}

def find_content_bbox_px(pil_image, white_threshold=245):
    mask = pil_image.convert("L").point(lambda p: 255 if p < white_threshold else 0)
    return mask.getbbox()

def make_borderless_pdf_and_png(raw_pdf, final_pdf, final_png, png_dpi=200):
    if fitz is not None:
        src = fitz.open(raw_pdf)
        page = src[0]

        probe_dpi = 200
        pix = page.get_pixmap(dpi=probe_dpi, alpha=False)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        bbox = find_content_bbox_px(img)

        if bbox:
            k = 72.0 / probe_dpi
            inset = 0.4
            clip = fitz.Rect(bbox[0] * k + inset, bbox[1] * k + inset,
                             bbox[2] * k - inset, bbox[3] * k - inset) & page.rect
        else:
            clip = page.rect

        out = fitz.open()
        target_w = A4_WIDTH_PT
        target_h = target_w * clip.height / clip.width
        new_page = out.new_page(width=target_w, height=target_h)
        new_page.show_pdf_page(new_page.rect, src, 0, clip=clip)
        out.save(final_pdf, garbage=3, deflate=True)
        out.close()
        src.close()

        chk = fitz.open(final_pdf)
        chk[0].get_pixmap(dpi=png_dpi, alpha=False).save(final_png)
        chk.close()
        return

    try:
        from pdf2image import convert_from_path
    except ImportError:
        os.replace(raw_pdf, final_pdf)
        raise RuntimeError("สร้างรูปภาพไม่ได้ ให้ติดตั้ง PyMuPDF ด้วยคำสั่ง: pip install pymupdf")

    os.replace(raw_pdf, final_pdf)
    pil_img = convert_from_path(final_pdf, dpi=png_dpi)[0].convert("RGB")
    bbox = find_content_bbox_px(pil_img)
    if bbox:
        pil_img = pil_img.crop(bbox)
    pil_img.save(final_png, "PNG")


BOT_VERSION = "v4-borderless-check-fixed"
CONVERT_STATUS = {}

def status_line(doc_number):
    s = CONVERT_STATUS.get(doc_number)
    return f"\n\n🔧 {s}" if s else ""


# =========================================================
# 📸 ฟังก์ชันรับและจัดการรูปภาพที่ผู้ใช้อัปโหลดเข้ามาในแชท
# =========================================================
async def handle_photo_upload(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        photo_file = await update.message.photo[-1].get_file()
        os.makedirs("downloads", exist_ok=True)
        file_path = f"downloads/photo_{update.message.from_user.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.jpg"
        await photo_file.download_to_drive(file_path)
        await update.message.reply_text("✅ บันทึกรูปภาพเรียบร้อยแล้วครับ!", reply_markup=main_menu_markup)
    except Exception as e:
        print(f"Photo upload error: {e}")
        await update.message.reply_text("❌ อัปโหลดรูปไม่สำเร็จ กรุณาลองใหม่อีกครั้ง", reply_markup=main_menu_markup)


def list_registry_printers():
    result = []
    try:
        import winreg
        key = winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                             r"Software\Microsoft\Windows NT\CurrentVersion\Devices")
        i = 0
        while True:
            try:
                name, data, _ = winreg.EnumValue(key, i)
            except OSError:
                break
            i += 1
            result.append((name, str(data)))
    except Exception:
        pass
    return result

def find_pdf_printer_name_for_excel():
    for name, data in list_registry_printers():
        if "microsoft print to pdf" in name.lower():
            port = data.split(",")[-1].strip()
            return f"{name} on {port}"
    return None

def use_pdf_printer(excel):
    candidates = []
    name = find_pdf_printer_name_for_excel()
    if name:
        candidates.append(name)
    candidates += [f"Microsoft Print to PDF on Ne{n:02d}:" for n in range(0, 100)]
    for cand in candidates:
        try:
            excel.ActivePrinter = cand
            return cand
        except Exception:
            continue
    return None

def diagnose_environment():
    lines = [f"🔧 เวอร์ชันบอท: {BOT_VERSION}"]
    lines.append("PyMuPDF: " + ("✅ ติดตั้งแล้ว" if fitz is not None else "❌ ยังไม่ได้ติดตั้ง"))
    try:
        import pdf2image  # noqa: F401
        lines.append("pdf2image: ✅")
    except ImportError:
        lines.append("pdf2image: -")

    printers = list_registry_printers()
    if printers:
        lines.append("เครื่องพิมพ์ที่เห็น: " + " | ".join(f"{n}" for n, _ in printers))
    else:
        lines.append("เครื่องพิมพ์ที่เห็น: (อ่านรายการไม่ได้)")
    lines.append("พบ Microsoft Print to PDF: " + (find_pdf_printer_name_for_excel() or "❌ ไม่พบ"))
    return lines

def convert_single_sheet_to_pdf_and_image(excel_filename, target_sheet_name, doc_number):
    pdf_filename = f"{doc_number}.pdf"
    image_filename = f"{doc_number}.png"
    raw_pdf = f"raw_{doc_number}.pdf"
    pdf_created, img_created = False, False
    CONVERT_ERRORS.pop(doc_number, None)

    excel = None
    wb = None
    used_printer = None
    try:
        import win32com.client
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False

        abs_excel = os.path.abspath(excel_filename)
        abs_raw_pdf = os.path.abspath(raw_pdf)

        wb = excel.Workbooks.Open(abs_excel)

        target_ws = None
        for ws in wb.Worksheets:
            if ws.Name == target_sheet_name:
                target_ws = ws
                break

        if target_ws is None:
            raise RuntimeError(f"ไม่พบชีทชื่อ '{target_sheet_name}' ในไฟล์ Excel")

        used_printer = use_pdf_printer(excel)

        target_ws.PageSetup.PrintArea = target_ws.UsedRange.Address
        target_ws.PageSetup.LeftMargin = 0
        target_ws.PageSetup.RightMargin = 0
        target_ws.PageSetup.TopMargin = 0
        target_ws.PageSetup.BottomMargin = 0
        target_ws.PageSetup.HeaderMargin = 0
        target_ws.PageSetup.FooterMargin = 0
        target_ws.PageSetup.Orientation = 1
        target_ws.PageSetup.PaperSize = 9
        target_ws.PageSetup.Zoom = False
        target_ws.PageSetup.FitToPagesWide = 1
        target_ws.PageSetup.FitToPagesTall = 1
        target_ws.PageSetup.CenterHorizontally = True
        target_ws.PageSetup.CenterVertically = False

        target_ws.ExportAsFixedFormat(0, abs_raw_pdf)
        pdf_created = True
    except Exception as e:
        print(f"Excel COM PDF Error: {e}")
        CONVERT_ERRORS[doc_number] = f"ส่งออก PDF จาก Excel ไม่สำเร็จ: {e}"
    finally:
        try:
            if wb is not None:
                wb.Close(False)
        except Exception:
            pass
        try:
            if excel is not None:
                excel.Quit()
        except Exception:
            pass

    CONVERT_STATUS[doc_number] = f"{BOT_VERSION} | เครื่องพิมพ์: {used_printer or 'สลับเป็น PDF ไม่สำเร็จ'}"

    if pdf_created and os.path.exists(raw_pdf):
        try:
            make_borderless_pdf_and_png(raw_pdf, pdf_filename, image_filename)
            img_created = os.path.exists(image_filename)
        except Exception as e:
            print(f"Borderless/PNG Error: {e}")
            CONVERT_ERRORS[doc_number] = f"ตัดขอบ/สร้างรูปภาพไม่สำเร็จ: {e}"
            if not os.path.exists(pdf_filename):
                try:
                    os.replace(raw_pdf, pdf_filename)
                except Exception:
                    pass
        finally:
            if os.path.exists(raw_pdf):
                try:
                    os.remove(raw_pdf)
                except Exception:
                    pass

    pdf_ok = os.path.exists(pdf_filename)
    return (pdf_filename if pdf_ok else None), (image_filename if img_created else None)

# ---------------------------------------------------------
# เมนูหลัก
# ---------------------------------------------------------
main_menu_markup = ReplyKeyboardMarkup(
    [
        ["📄 ออกใบเสนอราคา", "📑 ออกอินวอย (Invoice)"],
        ["🧾 ออกใบเสร็จรับเงิน", "📋 รายชื่อลูกค้าที่บันทึกไว้"],
        ["📖 คู่มือการใช้งาน"]
    ],
    resize_keyboard=True
)

# =========================================================
# 1. ออกใบเสนอราคา (Quotation)
# =========================================================
QUOT_FORM, INPUT_FIELD = range(2)

def build_quotation_keyboard(data):
    keyboard = [
        [InlineKeyboardButton(f"👤 ชื่อลูกค้า: {data.get('customer_name')}", callback_data="edit_name")],
        [InlineKeyboardButton(f"📍 ที่อยู่ 1: {data.get('address_line1')}", callback_data="edit_addr1")],
        [InlineKeyboardButton(f"📍 ที่อยู่ 2: {data.get('address_line2')}", callback_data="edit_addr2")],
        [
            InlineKeyboardButton(f"🆔 เลขภาษี: {data.get('tax_id')}", callback_data="edit_tax"),
            InlineKeyboardButton(f"📞 โทร: {data.get('phone')}", callback_data="edit_phone")
        ],
        [InlineKeyboardButton(f"📦 สินค้า: {data.get('item_name')}", callback_data="edit_item")],
        [
            InlineKeyboardButton(f"🔢 จำนวน: {data.get('qty'):,.0f}", callback_data="edit_qty"),
            InlineKeyboardButton(f"💰 ราคา/หน่วย: {data.get('price_per_unit'):,.2f}", callback_data="edit_price")
        ],
        [InlineKeyboardButton(f"📝 หมายเหตุ: {data.get('remarks')}", callback_data="edit_remarks")],
        [InlineKeyboardButton("🚀 สร้างใบเสนอราคา (เลือกรูปแบบไฟล์)", callback_data="generate_now")],
        [InlineKeyboardButton("❌ ยกเลิก", callback_data="cancel_menu")]
    ]
    return InlineKeyboardMarkup(keyboard)

async def start_quotation(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data['quot_data'] = {
        'customer_name': 'บริษัท คลีน แอนด์ แคร์ พูล เซอร์วิส จำกัด',
        'address_line1': '162/101 หมู่ที่ 9 ตำบลสวนหลวง',
        'address_line2': 'อำเภอกระทุ่มแบน จังหวัดสมุทรสาคร 74110',
        'tax_id': '0745564002745',
        'phone': '0891351884',
        'item_name': 'เตียงผู้ป่วย 3 ไกร์ระบบมือหมุน',
        'qty': 30.0,
        'price_per_unit': 18300.0,
        'remarks': '-'
    }
    await update.message.reply_text(
        "📄 **ระบบออกใบเสนอราคา (ไร้ขอบขาว)**\nกดปุ่มแก้ไขข้อมูล หรือกดสร้างเอกสารได้เลยครับ:",
        reply_markup=build_quotation_keyboard(context.user_data['quot_data'])
    )
    return QUOT_FORM

async def handle_quotation_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    action = query.data

    if action == "cancel_menu":
        await query.message.edit_text("❌ ยกเลิกรายการแล้ว")
        return ConversationHandler.END

    if action == "generate_now":
        await query.message.edit_text("⏳ กำลังสร้างไฟล์ PDF และรูปภาพ...")
        data = context.user_data['quot_data']
        doc_number = get_next_doc_number()
        date_str, due_str, date_iso = get_thai_dates()

        cust_info = {
            "customer_name": data['customer_name'],
            "address_line1": data['address_line1'],
            "address_line2": data['address_line2'],
            "tax_id": data['tax_id'],
            "phone": data['phone']
        }
        save_or_update_customer(cust_info)

        db = load_db()
        db["documents"].append({**data, "doc_number": doc_number, "date_iso": date_iso, "type": "quotation"})
        save_db(db)

        wb = openpyxl.load_workbook("บริษัท วินเนอร์ เชียงราย จำกัด.xlsx")
        ws = wb['เสนอราคา']
        ws['C9'] = f"      {data['customer_name']}"
        ws['C10'] = f"      {data['address_line1']}"
        ws['C11'] = f"      {data['address_line2']}"
        ws['B12'] = f"เลขกำกับภาษี    {data['tax_id']}"
        ws['E12'] = f"เบอร์โทรศัพท์  {data['phone']}"
        ws['H9'] = doc_number
        ws['H10'] = date_str
        ws['H11'] = due_str
        ws['C21'] = data['item_name']
        ws['F21'] = data['qty']
        ws['G21'] = data['price_per_unit']
        if data['remarks'] != "-":
            ws['B33'] = f"หมายเหตุ : {data['remarks']}"

        temp_excel = f"temp_{doc_number}.xlsx"
        wb.save(temp_excel)
        
        pdf_file, img_file = convert_single_sheet_to_pdf_and_image(temp_excel, 'เสนอราคา', doc_number)
        if os.path.exists(temp_excel): os.remove(temp_excel)

        context.bot_data[f"pdf_{doc_number}"] = pdf_file
        context.bot_data[f"img_{doc_number}"] = img_file

        file_select_keyboard = [
            [
                InlineKeyboardButton("📄 รับไฟล์ PDF", callback_data=f"getpdf_{doc_number}"),
                InlineKeyboardButton("🖼️ รับเป็นรูปภาพ", callback_data=f"getimg_{doc_number}")
            ]
        ]
        await query.message.reply_text(
            f"✅ **สร้างใบเสนอราคาสำเร็จ!** (เลขที่: `{doc_number}`)\nเลือกรูปแบบไฟล์ที่ต้องการด้านล่าง:{status_line(doc_number)}",
            reply_markup=InlineKeyboardMarkup(file_select_keyboard)
        )
        return ConversationHandler.END

    field_map = {
        "edit_name": ("customer_name", "กรุณาพิมพ์ชื่อลูกค้า:"),
        "edit_addr1": ("address_line1", "กรุณาพิมพ์ที่อยู่ บรรทัดที่ 1:"),
        "edit_addr2": ("address_line2", "กรุณาพิมพ์ที่อยู่ บรรทัดที่ 2:"),
        "edit_tax": ("tax_id", "กรุณาพิมพ์เลขประจำตัวผู้เสียภาษี:"),
        "edit_phone": ("phone", "กรุณาพิมพ์เบอร์โทรศัพท์:"),
        "edit_item": ("item_name", "กรุณาพิมพ์ชื่อรายการสินค้า:"),
        "edit_qty": ("qty", "กรุณาพิมพ์จำนวนสินค้า:"),
        "edit_price": ("price_per_unit", "กรุณาพิมพ์ราคาต่อหน่วย:"),
        "edit_remarks": ("remarks", "กรุณาพิมพ์หมายเหตุ:")
    }

    if action in field_map:
        context.user_data['editing_field'] = field_map[action][0]
        await query.message.reply_text(field_map[action][1])
        return INPUT_FIELD

    return QUOT_FORM

async def receive_quotation_input(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    field = context.user_data.get('editing_field')
    if field:
        if field in ['qty', 'price_per_unit']:
            try: context.user_data['quot_data'][field] = float(text)
            except: pass
        else:
            context.user_data['quot_data'][field] = text

    await update.message.reply_text(
        "🔄 อัปเดตข้อมูลเรียบร้อยแล้ว:",
        reply_markup=build_quotation_keyboard(context.user_data['quot_data'])
    )
    return QUOT_FORM


# =========================================================
# 2. ออกอินวอย (Invoice)
# =========================================================
INV_SELECT, INV_VAL = range(2)

async def start_invoice(update: Update, context: ContextTypes.DEFAULT_TYPE):
    db = load_db()
    docs = db.get("documents", [])[-5:]

    if not docs:
        await update.message.reply_text("⚠️ ยังไม่มีประวัติเอกสาร กรุณาออกใบเสนอราคาก่อนครับ", reply_markup=main_menu_markup)
        return ConversationHandler.END

    keyboard = []
    for i, doc in enumerate(docs):
        keyboard.append([InlineKeyboardButton(f"📁 {doc.get('customer_name')} ({doc.get('doc_number')})", callback_data=f"sel_inv_{i}")])
    
    context.user_data['available_docs'] = docs
    await update.message.reply_text(
        "📑 **ออกอินวอย (Invoice)**\nเลือกรายชื่อจากประวัติเพื่อดึงข้อมูลมาออกอินวอย:",
        reply_markup=InlineKeyboardMarkup(keyboard)
    )
    return INV_SELECT

async def invoice_doc_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    idx = int(query.data.split("_")[-1])
    doc = context.user_data['available_docs'][idx]
    context.user_data['inv_data'] = doc

    await query.message.reply_text(
        f"✅ ดึงข้อมูลของ: **{doc['customer_name']}** เรียบร้อย\n\n"
        "พิมพ์ยอดชำระ (เช่นพิมพ์ `60` สำหรับ 60% หรือระบุจำนวนเงินบาท เช่น `50000`):"
    )
    return INV_VAL

async def invoice_generate(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    data = context.user_data['inv_data']
    doc_number = get_next_doc_number()
    date_str, _, _ = get_thai_dates()
    
    total_val = data['qty'] * data['price_per_unit'] * 1.07
    try:
        val = float(text)
        if val <= 100 and "." not in text:
            paid_amt = total_val * (val / 100)
            pct = int(val)
        else:
            paid_amt = val
            pct = round((paid_amt / total_val) * 100)
        remark_text = f"**ลูกค้าชำระยอด {pct}% จากยอดเต็มเป็นจำนวนเงิน {paid_amt:,.2f} บาท**"
    except:
        paid_amt = total_val * 0.60
        remark_text = "**ลูกค้าชำระยอด 60% จากยอดเต็ม**"

    wb = openpyxl.load_workbook("บริษัท วินเนอร์ เชียงราย จำกัด.xlsx")
    ws = wb['อินวอย(ช)']
    ws['B12'] = data['customer_name']
    ws['C13'] = data['address_line1']
    ws['C14'] = data['address_line2']
    ws['D15'] = data['tax_id']
    ws['D16'] = data['phone']

    ws['I13'] = doc_number
    ws['I14'] = date_str
    ws['C20'] = data['item_name']
    ws['G20'] = data['qty']
    ws['H20'] = data['price_per_unit']
    ws['B29'] = remark_text
    ws['J36'] = paid_amt

    temp_excel = f"temp_inv_{doc_number}.xlsx"
    wb.save(temp_excel)
    
    await update.message.reply_text("⏳ กำลังสร้างไฟล์ PDF และรูปภาพ...")
    pdf_file, img_file = convert_single_sheet_to_pdf_and_image(temp_excel, 'อินวอย(ช)', doc_number)
    if os.path.exists(temp_excel): os.remove(temp_excel)

    context.bot_data[f"pdf_{doc_number}"] = pdf_file
    context.bot_data[f"img_{doc_number}"] = img_file

    file_select_keyboard = [
        [
            InlineKeyboardButton("📄 รับไฟล์ PDF", callback_data=f"getpdf_{doc_number}"),
            InlineKeyboardButton("🖼️ รับเป็นรูปภาพ", callback_data=f"getimg_{doc_number}")
        ]
    ]
    await update.message.reply_text(
        f"✅ **สร้างใบอินวอยสำเร็จ!** (เลขที่: `{doc_number}`)\nเลือกรูปแบบไฟล์ที่ต้องการด้านล่าง:{status_line(doc_number)}",
        reply_markup=InlineKeyboardMarkup(file_select_keyboard)
    )
    return ConversationHandler.END


# =========================================================
# 3. ออกใบเสร็จรับเงิน (Receipt)
# =========================================================
REC_SELECT, REC_VAL = range(2)

async def start_receipt(update: Update, context: ContextTypes.DEFAULT_TYPE):
    db = load_db()
    docs = db.get("documents", [])[-5:]

    if not docs:
        await update.message.reply_text("⚠️ ยังไม่มีประวัติเอกสาร กรุณาออกใบเสนอราคาก่อนครับ", reply_markup=main_menu_markup)
        return ConversationHandler.END

    keyboard = []
    for i, doc in enumerate(docs):
        keyboard.append([InlineKeyboardButton(f"📁 {doc.get('customer_name')} ({doc.get('doc_number')})", callback_data=f"sel_rec_{i}")])
    
    context.user_data['available_docs'] = docs
    await update.message.reply_text(
        "🧾 **ออกใบเสร็จรับเงิน**\nเลือกรายชื่อจากประวัติเพื่อดึงข้อมูลมาออกใบเสร็จ:",
        reply_markup=InlineKeyboardMarkup(keyboard)
    )
    return REC_SELECT

async def receipt_doc_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    idx = int(query.data.split("_")[-1])
    doc = context.user_data['available_docs'][idx]
    context.user_data['rec_data'] = doc

    await query.message.reply_text(
        f"✅ ดึงข้อมูลของ: **{doc['customer_name']}** เรียบร้อย\n\n"
        "พิมพ์ยอดชำระเงินที่รับจริง (เช่นพิมพ์ `50000` บาท หรือพิมพ์ % เช่น `100`):"
    )
    return REC_VAL

async def receipt_generate(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    data = context.user_data['rec_data']
    doc_number = get_next_doc_number()
    date_str, due_str, _ = get_thai_dates()
    
    total_val = data['qty'] * data['price_per_unit'] * 1.07
    try:
        val = float(text)
        if val <= 100 and "." not in text:
            paid_amt = total_val * (val / 100)
            pct = int(val)
        else:
            paid_amt = val
            pct = round((paid_amt / total_val) * 100)
        remark_text = f"**ลูกค้าชำระยอด {pct}% จากยอดเต็มเป็นจำนวนเงิน {paid_amt:,.2f} บาท**"
    except:
        paid_amt = total_val
        pct = 100
        remark_text = "**ลูกค้าชำระยอดเต็ม 100%**"

    wb = openpyxl.load_workbook("บริษัท วินเนอร์ เชียงราย จำกัด.xlsx")
    ws = wb['รับเงิน']
    ws['C10'] = data['customer_name']
    ws['D11'] = data['address_line1']
    ws['D12'] = data['address_line2']
    ws['E13'] = data['tax_id']
    ws['G9'] = data['phone']

    ws['I10'] = f"เลขที่          {doc_number}"
    ws['I11'] = f"วันที่           {date_str}"
    ws['I12'] = f"ครบกำหนด   {due_str}"
    ws['D16'] = data['item_name']
    ws['G16'] = data['qty']
    ws['H16'] = data['price_per_unit']
    ws['C30'] = remark_text
    ws['I32'] = paid_amt

    temp_excel = f"temp_rec_{doc_number}.xlsx"
    wb.save(temp_excel)
    
    await update.message.reply_text("⏳ กำลังสร้างไฟล์ PDF และรูปภาพ...")
    pdf_file, img_file = convert_single_sheet_to_pdf_and_image(temp_excel, 'รับเงิน', doc_number)
    if os.path.exists(temp_excel): os.remove(temp_excel)

    context.bot_data[f"pdf_{doc_number}"] = pdf_file
    context.bot_data[f"img_{doc_number}"] = img_file

    file_select_keyboard = [
        [
            InlineKeyboardButton("📄 รับไฟล์ PDF", callback_data=f"getpdf_{doc_number}"),
            InlineKeyboardButton("🖼️ รับเป็นรูปภาพ", callback_data=f"getimg_{doc_number}")
        ]
    ]
    await update.message.reply_text(
        f"✅ **สร้างใบเสร็จรับเงินสำเร็จ!** (เลขที่: `{doc_number}`)\nเลือกรูปแบบไฟล์ที่ต้องการด้านล่าง:{status_line(doc_number)}",
        reply_markup=InlineKeyboardMarkup(file_select_keyboard)
    )
    return ConversationHandler.END


# =========================================================
# 4. ฟังก์ชันส่งไฟล์ PDF หรือ รูปภาพ ตามปุ่มที่กดเลือก
# =========================================================
async def handle_file_selection_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    data = query.data

    if data.startswith("getpdf_"):
        doc_num = data.replace("getpdf_", "")
        pdf_file = context.bot_data.get(f"pdf_{doc_num}")
        if pdf_file and os.path.exists(pdf_file):
            with open(pdf_file, 'rb') as f:
                await query.message.reply_document(document=f, caption=f"📄 ไฟล์ PDF เอกสารเลขที่ {doc_num}", reply_markup=main_menu_markup)
        else:
            err = CONVERT_ERRORS.get(doc_num, "ไม่ทราบสาเหตุ")
            await query.message.reply_text(f"⚠️ ไม่พบไฟล์ PDF\nสาเหตุ: {err}", reply_markup=main_menu_markup)

    elif data.startswith("getimg_"):
        doc_num = data.replace("getimg_", "")
        img_file = context.bot_data.get(f"img_{doc_num}")
        if img_file and os.path.exists(img_file):
            with open(img_file, 'rb') as f:
                await query.message.reply_photo(photo=f, caption=f"🖼️ รูปภาพเอกสารเลขที่ {doc_num}", reply_markup=main_menu_markup)
        else:
            err = CONVERT_ERRORS.get(doc_num, "ไม่ทราบสาเหตุ")
            await query.message.reply_text(f"⚠️ ไม่พบไฟล์รูปภาพ\nสาเหตุ: {err}", reply_markup=main_menu_markup)

# =========================================================
# 5. รายชื่อลูกค้า และ คู่มือการใช้งาน
# =========================================================
async def show_customers(update: Update, context: ContextTypes.DEFAULT_TYPE):
    db = load_db()
    customers = db.get("customers", [])
    if not customers:
        await update.message.reply_text("⚠️ ยังไม่มีรายชื่อลูกค้าในระบบ", reply_markup=main_menu_markup)
        return
    msg = f"📋 **รายชื่อลูกค้าทั้งหมด ({len(customers)} ราย):**\n\n"
    for idx, c in enumerate(customers, 1):
        msg += f"{idx}. **{c['customer_name']}**\n   📍 {c['address_line1']} {c['address_line2']}\n   🆔 เลขภาษี: {c['tax_id']} | 📞 โทร: {c['phone']}\n\n"
    await update.message.reply_text(msg, reply_markup=main_menu_markup)

async def show_guide(update: Update, context: ContextTypes.DEFAULT_TYPE):
    guide_text = (
        "📖 **คู่มือการใช้งานระบบออกเอกสาร**\n\n"
        "1️⃣ **เลือกเอกสาร:** สร้างใบเสนอราคา / อินวอย / ใบเสร็จ\n"
        "2️⃣ **เลือรูปแบบไฟล์:** หลังจากสร้างเสร็จจะมีปุ่มให้กดเลือกว่าจะรับเป็น **PDF** หรือ **รูปภาพ (PNG)**\n"
        "3️⃣ **ไร้ขอบขาว 100%:** เอกสารทุกไฟล์ถูกปรับขนาดพอดีขอบกระดาษเรียบร้อย\n"
        "💬 พิมพ์ /start เพื่อรีเซ็ตเมนูหลักได้ตลอดเวลาครับ"
    )
    await update.message.reply_text(guide_text, reply_markup=main_menu_markup)

async def check_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("🔍 กำลังตรวจสอบระบบ (อาจใช้เวลา 5-10 วินาที)...")
    lines = diagnose_environment()
    await update.message.reply_text("\n".join(lines), reply_markup=main_menu_markup)


async def start_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        f"🤖 **ระบบออกเอกสาร Winner Chiangrai** ({BOT_VERSION})\nเลือกเมนูด้านล่างเพื่อเริ่มใช้งานได้เลยครับ:",
        reply_markup=main_menu_markup
    )

async def cancel_process(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text("❌ ยกเลิกรายการเรียบร้อย", reply_markup=main_menu_markup)
    return ConversationHandler.END

# ---------------------------------------------------------
# Main App
# ---------------------------------------------------------
web_app = Flask('')
@web_app.route('/')
def home():
    return "Bot is running locally!"

def run_web():
    web_app.run(host='0.0.0.0', port=8080)

if __name__ == "__main__":
    init_db()
    if fitz is None:
        print("⚠️ ยังไม่ได้ติดตั้ง PyMuPDF — ปุ่ม 'รับเป็นรูปภาพ' อาจใช้ไม่ได้  รันคำสั่ง: pip install pymupdf")
    threading.Thread(target=run_web, daemon=True).start()
    
    request = HTTPXRequest(connect_timeout=30.0, read_timeout=30.0)
    app = ApplicationBuilder().token(TELEGRAM_BOT_TOKEN).request(request).build()
    
    quot_handler = ConversationHandler(
        entry_points=[MessageHandler(filters.Regex(r"^(📄 ออกใบเสนอราคา|ออกใบเสนอราคา)$"), start_quotation)],
        states={
            QUOT_FORM: [CallbackQueryHandler(handle_quotation_callback)],
            INPUT_FIELD: [MessageHandler(filters.TEXT & ~filters.COMMAND, receive_quotation_input)],
        },
        fallbacks=[CommandHandler("cancel", cancel_process)]
    )

    inv_handler = ConversationHandler(
        entry_points=[MessageHandler(filters.Regex(r"^(📑 ออกอินวอย \(Invoice\)|ออกอินวอย)$"), start_invoice)],
        states={
            INV_SELECT: [CallbackQueryHandler(invoice_doc_selected, pattern="^sel_inv_")],
            INV_VAL: [MessageHandler(filters.TEXT & ~filters.COMMAND, invoice_generate)],
        },
        fallbacks=[CommandHandler("cancel", cancel_process)]
    )

    rec_handler = ConversationHandler(
        entry_points=[MessageHandler(filters.Regex(r"^(🧾 ออกใบเสร็จรับเงิน|ออกใบเสร็จรับเงิน)$"), start_receipt)],
        states={
            REC_SELECT: [CallbackQueryHandler(receipt_doc_selected, pattern="^sel_rec_")],
            REC_VAL: [MessageHandler(filters.TEXT & ~filters.COMMAND, receipt_generate)],
        },
        fallbacks=[CommandHandler("cancel", cancel_process)]
    )

    app.add_handler(CommandHandler("start", start_command))
    app.add_handler(CommandHandler("check", check_command))
    app.add_handler(MessageHandler(filters.Regex(r"^(📋 รายชื่อลูกค้าที่บันทึกไว้)$"), show_customers))
    app.add_handler(MessageHandler(filters.Regex(r"^(📖 คู่มือการใช้งาน)$"), show_guide))
    
    # ✅ เพิ่มส่วนนี้เข้ามาเพื่อดักจับรูปภาพที่ผู้ใช้ส่งเข้ามา ป้องกันข้อผิดพลาด "อัปโหลดรูปไม่สำเร็จ"
    app.add_handler(MessageHandler(filters.PHOTO, handle_photo_upload))
    
    app.add_handler(quot_handler)
    app.add_handler(inv_handler)
    app.add_handler(rec_handler)
    app.add_handler(CallbackQueryHandler(handle_file_selection_callback, pattern="^(getpdf_|getimg_)"))
    
    print(f"เวอร์ชันบอท: {BOT_VERSION}")
    print("🤖 บอทพร้อมทำงานแล้ว...")
    app.run_polling()